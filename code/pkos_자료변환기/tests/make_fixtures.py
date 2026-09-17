# -*- coding: utf-8 -*-
"""
시험 자료 만들기 — 변환기가 까다로워하는 파일들을 일부러 만든다
=================================================================

실제 자료를 시험에 쓰면 개인정보가 섞이고, 다른 PC에는 그 자료가 없다.
그래서 **내용을 직접 지어서** 코드로 만든다. 어느 PC에서나 똑같이 만들어진다.

    python tests/make_fixtures.py <만들 폴더>

만드는 것 (각각 '무엇을 시험하려는지'가 이름에 드러나게 했다)

    이름과 속이 다른 파일   확장자 없는 파일      빈 파일
    표가 본문 중간에 있는 워드 · 글상자 · 그림    그룹 도형이 있는 PPT
    cp949 텍스트 · UTF-16 텍스트 · 쉼표 섞인 메모  탭 구분 표
    머리말(YAML)이 있는 마크다운                   div 로만 된 웹문서
    한글 이름이 cp949 로 적힌 ZIP (한국 윈도우 압축) 암호 걸린 PDF · 스캔 PDF
    같은 이름 다른 내용의 사진 두 장               찍은 날짜가 든 사진
    내용이 같은 사본                               이름이 겹치는 파일
    윈도우 잡동사니(desktop.ini · Thumbs.db · 바로가기)
    아주 깊은 폴더 (경로 260자 넘김)
    개인정보 함정 (주소 옆 카드번호 · 짧은 은행 이름 · '이름 설정')

옛 형식(.doc/.ppt/.xls)과 암호 걸린 오피스 파일은 코드로 만들 수 없어서
tests/fixtures/ 에 미리 만들어 둔 것을 복사한다 (tools/ 의 스크립트로 만들었다).
"""

from __future__ import annotations

import io
import os
import sys
import json
import shutil
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
STATIC = os.path.join(HERE, "fixtures")


# ─────────────────────────────────────────────────────────────
def _png(path: str, w: int = 220, h: int = 160, seed: int = 1):
    """압축이 안 되는 잡음 그림 — 8KB 를 확실히 넘겨 '아이콘'으로 걸러지지 않게"""
    from PIL import Image
    import random
    rnd = random.Random(seed)
    data = bytes(rnd.randrange(256) for _ in range(w * h * 3))
    Image.frombytes("RGB", (w, h), data).save(path, "PNG")


def _jpeg_with_exif(path: str, seed: int, when: str | None = None):
    from PIL import Image
    import random
    rnd = random.Random(seed)
    data = bytes(rnd.randrange(256) for _ in range(300 * 200 * 3))
    img = Image.frombytes("RGB", (300, 200), data)
    exif = img.getexif()
    exif[0x010F] = "Apple"
    exif[0x0110] = "iPhone 12"
    if when:
        exif.get_ifd(0x8769)[0x9003] = when
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    img.save(path, "JPEG", exif=exif, quality=90)


def _write(path: str, data, mode: str = "w", encoding: str = "utf-8"):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    if "b" in mode:
        with open(path, mode) as f:
            f.write(data)
    else:
        with io.open(path, mode, encoding=encoding, newline="") as f:
            f.write(data)


# ─────────────────────────────────────────────────────────────
def make_docx(path: str, img: str, title: str = "워드 시험 문서"):
    import docx
    from docx.oxml import parse_xml
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "머리글에만 있는 글"
    d.add_heading(title, level=1)
    d.add_paragraph("표 앞 문단입니다.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "항목", "값"
    t.cell(1, 0).text, t.cell(1, 1).text = "사과", "3"
    d.add_paragraph("표 뒤 문단입니다.")
    d.add_heading("둘째 절", level=2)
    d.add_picture(img)
    d.add_paragraph("그림 뒤 문단입니다.")
    # VML 글상자 — 워드에서 흔히 쓰는 '텍스트 상자'
    box = parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>'
        '<v:shape style="width:200pt;height:40pt"><v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t>글상자 안의 글</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
    d.element.body.insert(len(d.element.body) - 1, box)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    d.save(path)


def make_pptx(path: str, img: str, title: str = "발표 시험 자료"):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    p = Presentation()
    s1 = p.slides.add_slide(p.slide_layouts[0])
    s1.shapes.title.text = title
    s1.placeholders[1].text = "부제목입니다"

    s2 = p.slides.add_slide(p.slide_layouts[5])          # 제목만
    s2.shapes.title.text = "그룹 도형 슬라이드"
    grp = s2.shapes.add_group_shape()
    for i, y in enumerate((2, 3), 1):
        tb = grp.shapes.add_textbox(Inches(1), Inches(y), Inches(4), Inches(0.6))
        tb.text_frame.text = f"그룹 안 글상자 {i}"
    s2.notes_slide.notes_text_frame.text = "발표자 노트 내용입니다"

    s3 = p.slides.add_slide(p.slide_layouts[5])
    s3.shapes.title.text = "표와 그림 슬라이드"
    tbl = s3.shapes.add_table(2, 2, Inches(0.5), Inches(1.5), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "이름", "점수"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "가", "90"
    s3.shapes.add_picture(img, Inches(5), Inches(1.5), Inches(3))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    p.save(path)


def make_xlsx(path: str, subject: str = "국어"):
    import datetime
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "성적"
    ws.append(["번호", "과목", "점수", "날짜"])
    ws.append([1, subject, 95, datetime.date(2026, 3, 2)])
    ws.append([2, "수학", 87.5, datetime.date(2026, 3, 3)])
    wb.create_sheet("빈시트")
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def make_pdfs(folder: str, img: str):
    import pymupdf
    # 글자가 있는 PDF (2쪽, 둘째 쪽에 그림)
    d = pymupdf.open()
    pg = d.new_page()
    pg.insert_text((72, 90), "PDF 첫째 쪽 본문입니다", fontname="korea", fontsize=14)
    pg = d.new_page()
    pg.insert_text((72, 90), "PDF 둘째 쪽 본문입니다", fontname="korea", fontsize=14)
    pg.insert_image(pymupdf.Rect(72, 120, 292, 280), filename=img)
    d.save(os.path.join(folder, "글자PDF.pdf"))
    d.close()

    # 암호 걸린 PDF
    d = pymupdf.open()
    d.new_page().insert_text((72, 90), "비밀 문서", fontname="korea")
    d.save(os.path.join(folder, "암호PDF.pdf"),
           encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1234", owner_pw="1234")
    d.close()

    # 스캔한 PDF — 글자가 그림으로만 들어 있다
    d = pymupdf.open()
    pg = d.new_page()
    pg.insert_image(pymupdf.Rect(50, 50, 500, 400), filename=img)
    d.save(os.path.join(folder, "스캔PDF.pdf"))
    d.close()


def make_hwpx(path: str, img: str):
    """최소한의 한글 hwpx. 기존 엔진이 틀리던 네 가지를 일부러 담는다.
        · 서식이 바뀌어 낱말 중간에서 run 이 나뉨   ('안' + '녕하세요')
        · <hp:t> 안의 탭 요소 뒤에 오는 글
        · 표 칸 안의 표 (안쪽 칸 주소가 바깥 칸을 덮어쓰면 안 된다)
        · 구역(section) 11개 — 글자 순 정렬이면 section10 이 section2 앞에 온다
    """
    ns = ('xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
          'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
          'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core"')

    def cell(r, c, inner):
        return (f'<hp:tc><hp:subList>{inner}</hp:subList>'
                f'<hp:cellAddr colAddr="{c}" rowAddr="{r}"/></hp:tc>')

    def para(text):
        return f'<hp:p><hp:run><hp:t>{text}</hp:t></hp:run></hp:p>'

    nested = ('<hp:p><hp:run><hp:tbl rowCnt="1" colCnt="2"><hp:tr>'
              + cell(0, 0, para("속1")) + cell(0, 1, para("속2")) +
              '</hp:tr></hp:tbl></hp:run></hp:p>')
    first = (f'<?xml version="1.0" encoding="UTF-8"?><hs:sec {ns}>'
             '<hp:p><hp:run><hp:t>안</hp:t></hp:run><hp:run><hp:t>녕하세요</hp:t></hp:run></hp:p>'
             '<hp:p><hp:run><hp:t>탭<hp:tab/>뒤의 글</hp:t></hp:run></hp:p>'
             '<hp:p><hp:run><hp:tbl rowCnt="2" colCnt="2">'
             '<hp:tr>' + cell(0, 0, para("가")) + cell(0, 1, para("나")) + '</hp:tr>'
             '<hp:tr>' + cell(1, 0, para("다")) + cell(1, 1, nested) + '</hp:tr>'
             '</hp:tbl></hp:run></hp:p>'
             '<hp:p><hp:run><hp:pic><hc:img binaryItemIDRef="image1"/></hp:pic></hp:run></hp:p>'
             + para("구역 0") + '</hs:sec>')
    with open(img, "rb") as f:
        png = f.read()
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("version.xml", '<?xml version="1.0"?><hv:HCFVersion xmlns:hv="x"/>')
        z.writestr("Contents/content.hpf",
                   '<?xml version="1.0"?><opf:package xmlns:opf="http://www.idpf.org/2007/opf/">'
                   '<opf:manifest><opf:item id="image1" href="BinData/image1.png" '
                   'media-type="image/png"/></opf:manifest></opf:package>')
        z.writestr("Contents/section0.xml", first)
        for i in range(1, 11):
            z.writestr(f"Contents/section{i}.xml",
                       f'<?xml version="1.0" encoding="UTF-8"?><hs:sec {ns}>{para(f"구역 {i}")}</hs:sec>')
        z.writestr("BinData/image1.png", png)


def make_odt(path: str, img: str):
    """한글에서 오픈도큐먼트로 내보낸 모양 — 표가 '문단 안의 틀 속 글상자'에 들어 있다.
    <text:s text:c="3"/> 는 공백 세 칸이다."""
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
          'xmlns:xlink="http://www.w3.org/1999/xlink"')
    content = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns}>'
        '<office:body><office:text>'
        '<text:h text:outline-level="1">공문 제목</text:h>'
        '<text:p>A<text:s text:c="3"/>B</text:p>'
        '<text:p><draw:frame><draw:text-box>'
        '<table:table><table:table-row>'
        '<table:table-cell><text:p>항목</text:p></table:table-cell>'
        '<table:table-cell><text:p>내용</text:p></table:table-cell></table:table-row>'
        '<table:table-row><table:table-cell table:number-columns-spanned="2">'
        '<text:p>합친 칸</text:p></table:table-cell><table:covered-table-cell/></table:table-row>'
        '<table:table-row><table:table-cell><text:p><text:span>'
        '<draw:frame><draw:image xlink:href="Pictures/칸속그림.png"/></draw:frame>'
        '</text:span></text:p></table:table-cell>'
        '<table:table-cell><text:p><draw:frame><draw:text-box>'
        '<text:p>칸 속 글상자 글</text:p></draw:text-box></draw:frame></text:p>'
        '</table:table-cell></table:table-row>'
        '</table:table></draw:text-box></draw:frame></text:p>'
        '<text:p><draw:frame><draw:image xlink:href="Pictures/그림1.png"/></draw:frame></text:p>'
        '<text:p>끝 문단</text:p>'
        '</office:text></office:body></office:document-content>')
    with open(img, "rb") as f:
        png = f.read()
    cell_img = os.path.join(os.path.dirname(img), "칸속그림.png")
    _png(cell_img, seed=31)                                   # 내용이 다른 진짜 그림
    with open(cell_img, "rb") as f:
        png2 = f.read()
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
        z.writestr("Pictures/그림1.png", png)
        z.writestr("Pictures/칸속그림.png", png2)


def patch_content_type(src: str, dst: str, old: bytes, new: bytes):
    """서식파일(.dotx)·쇼(.ppsx) 는 속이 같고 '본체 종류' 표시만 다르다."""
    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                assert old in data, "본체 종류 표시를 찾지 못함"
                data = data.replace(old, new)
            zout.writestr(item, data)


def make_blog_pdf(path: str, posts: list[tuple[str, str, str, list[str]]], img: str):
    """네이버 블로그 '인쇄 → PDF 로 저장' 백업과 같은 모양.
    글머리는 '날짜시각' 줄 바로 다음 'blog.naver.com/아이디/글번호' 줄이고,
    쪽마다 바닥글 'N · 블로그이름' 이 붙는다. 긴 글은 다음 쪽으로 넘어간다."""
    import pymupdf
    d = pymupdf.open()
    page, y, pno = None, 0, 0

    def new_page():
        nonlocal page, y, pno
        if page is not None:
            page.insert_text((72, 810), f"{pno} · 시험블로그", fontname="korea", fontsize=8)
        page, y = d.new_page(), 60
        pno += 1

    def line(text):
        nonlocal y
        if page is None or y > 760:
            new_page()
        page.insert_text((72, y), text, fontname="korea", fontsize=11)
        y += 20

    for when, postid, title, body in posts:
        new_page()                                 # 글마다 새 쪽에서 시작
        line(when)
        line(f"http://blog.naver.com/testblog/{postid}")
        line(title)
        line("수업 이야기")
        for b in body:
            line(b)
        if img:
            page.insert_image(pymupdf.Rect(72, y, 272, y + 140), filename=img)
            y += 150
    new_page()
    d.delete_page(-1)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    d.save(path)
    d.close()


def make_cp949_zip(path: str):
    """한국 윈도우 압축 프로그램이 만든 것처럼, 파일 이름을 cp949 로 적고
    UTF-8 표시(0x800)는 켜지 않은 ZIP."""

    class CP949Info(zipfile.ZipInfo):
        def _encodeFilenameFlags(self):
            return self.filename.encode("cp949"), self.flag_bits & ~0x800

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in (("압축속_한글문서.txt", "압축 안에 든 한글 문서입니다.\n"),
                           ("하위/압축속_메모.md", "# 압축 속 메모\n\n내용입니다.\n")):
            zi = CP949Info(name)
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, text.encode("utf-8"))


# ─────────────────────────────────────────────────────────────
def build(root: str) -> str:
    """root 아래에 시험 자료를 만들고 root 를 돌려준다."""
    if os.path.exists(root):
        shutil.rmtree(root)
    os.makedirs(root)
    tmp = os.path.join(root, "_재료")
    os.makedirs(tmp)
    img = os.path.join(tmp, "그림.png")
    _png(img, seed=7)

    # ── 오피스 (새 형식)
    make_docx(os.path.join(root, "워드문서.docx"), img)
    make_pptx(os.path.join(root, "발표자료.pptx"), img)
    make_xlsx(os.path.join(root, "성적표.xlsx"))

    # 이름과 속이 다른 것 / 확장자 없는 것
    make_docx(os.path.join(root, "사실은워드.pdf"), img, "이름은 PDF 인 워드")
    make_pptx(os.path.join(root, "확장자없는발표"), img, "확장자 없는 발표")

    # 같은 제목 · 다른 내용 → 그림 폴더가 겹치면 안 된다
    _png(os.path.join(tmp, "그림2.png"), seed=99)
    make_docx(os.path.join(root, "같은제목", "보고서.docx"), img, "같은 제목")
    make_pptx(os.path.join(root, "같은제목", "보고서.pptx"),
              os.path.join(tmp, "그림2.png"), "같은 제목")

    # 서식파일·쇼파일 — python-docx/pptx 가 '본체 종류'만 보고 거부하던 것
    patch_content_type(
        os.path.join(root, "워드문서.docx"), os.path.join(root, "서식", "서식파일.dotx"),
        b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
    patch_content_type(
        os.path.join(root, "발표자료.pptx"), os.path.join(root, "서식", "쇼파일.ppsx"),
        b"presentationml.presentation.main+xml", b"presentationml.slideshow.main+xml")
    make_pptx(os.path.join(tmp, "매크로원본.pptx"), img, "매크로가 든 발표")
    patch_content_type(
        os.path.join(tmp, "매크로원본.pptx"), os.path.join(root, "서식", "매크로발표.pptm"),
        b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml")
    # 이름이 .xlsx 가 아닌 엑셀 — openpyxl 은 확장자만 보고 거부한다
    # (성적표.xlsx 를 복사하면 사본으로 걸러지므로 내용을 달리한다)
    make_xlsx(os.path.join(root, "서식", "엑셀인데이름이틀림.dat"), subject="영어")

    # 한글 hwpx · 오픈도큐먼트
    make_hwpx(os.path.join(root, "한글", "한글문서.hwpx"), img)
    make_odt(os.path.join(root, "한글", "공문.odt"), img)

    # 사용자가 만든 'assets' 폴더 — 결과 폴더 이름과 같아도 건너뛰면 안 된다
    _write(os.path.join(root, "assets", "내자료.txt"), "사용자가 만든 assets 폴더 안의 글입니다.\n")

    # ── 블로그 백업 PDF 두 개 — 글 2 가 양쪽에 겹친다
    long_body = [f"첫 글의 {i}번째 줄입니다. 쪽을 넘길 만큼 깁니다." for i in range(1, 70)]
    post1 = ("2018/04/15 19:47", "220000000001", '드론 수업 "첫날" 기록', long_body)
    post2 = ("2018/05/01 10:00", "220000000002", "둘째 글", ["문의는 010-9876-5432 로 주세요."])
    post3 = ("2018/06/10 08:30", "220000000003", "셋째 글(평등과 형평)", ["셋째 글 본문입니다."])
    make_blog_pdf(os.path.join(root, "블로그백업", "1_2.pdf"), [post1, post2], img)
    make_blog_pdf(os.path.join(root, "블로그백업", "2_3.pdf"), [post2, post3], img)

    # ── PDF
    make_pdfs(root, img)

    # ── 글자파일
    _write(os.path.join(root, "글", "cp949메모.txt"),
           "옛 윈도우 메모장으로 쓴 글입니다.\n둘째 줄입니다.\n", encoding="cp949")
    _write(os.path.join(root, "글", "유니코드메모.txt"),
           "\ufeffUTF-16 으로 저장한 메모입니다.\n", encoding="utf-16-le")
    _write(os.path.join(root, "글", "쉼표메모.txt"),
           "안녕하세요, 반갑습니다.\n오늘은, 좋은 날입니다.\n")
    _write(os.path.join(root, "글", "명단.tsv"),
           "번호\t이름\t반\n1\t가나다\t3\n2\t라마바\t4\n")
    _write(os.path.join(root, "글", "점수.csv"),
           "과목,점수\n국어,90\n", encoding="cp949")
    _write(os.path.join(root, "글", "노트.md"),
           "---\ntitle: 옵시디언 노트\ntags: [수업, 파이썬]\naliases:\n  - 노트\n---\n\n"
           "# 옵시디언 노트\n\n본문입니다.\n")
    _write(os.path.join(root, "글", "웹페이지.html"),
           '<html><head><meta charset="euc-kr"><title>웹 제목</title></head><body>'
           "<div>div 로만 된 첫 글</div><div><span>span 안의 글</span></div>"
           "<table><tr><th>가</th><th>나</th></tr><tr><td>1</td><td>2</td></tr></table>"
           "<pre>print('코드')</pre></body></html>", encoding="cp949")
    _write(os.path.join(root, "글", "강의자막.smi"),
           "<SAMI><BODY><SYNC Start=1000><P Class=KRCC>첫 번째 자막입니다.\n"
           "<SYNC Start=3000><P Class=KRCC>&nbsp;\n"
           "<SYNC Start=4000><P Class=KRCC>두 번째 자막입니다.\n</BODY></SAMI>\n", encoding="cp949")
    _write(os.path.join(root, "글", "영상자막.vtt"),
           "WEBVTT\n\n00:00:01.000 --> 00:00:03.000\n<i>웹 자막</i> 한 줄입니다.\n")
    _write(os.path.join(root, "글", "a.b"), "점이 들어간 이름\n")
    _write(os.path.join(root, "글", "a_b"), "밑줄이 들어간 이름\n")
    _write(os.path.join(root, "글", "빈파일.txt"), "")
    # 이름이 쓸모없으면 본문 첫 줄이 제목이 된다. 그 줄에 \ 와 " 가 있으면
    # 머리말(YAML)이 깨지기 쉽다. 파일 이름에는 \ 를 넣을 수 없으니 본문으로 시험한다.
    _write(os.path.join(root, "글", "문서1.txt"),
           'C:\\Users\\"따옴표" 경로 설명\n본문입니다.\n')

    # ── 사진 — 같은 이름, 다른 내용, 다른 폴더
    _jpeg_with_exif(os.path.join(root, "사진", "가", "IMG_0001.jpg"), 1,
                    "2019:07:21 14:03:22")
    _jpeg_with_exif(os.path.join(root, "사진", "나", "IMG_0001.jpg"), 2,
                    "2020:01:05 09:00:00")
    _write(os.path.join(root, "사진", "도식.svg"),
           '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="60">'
           '<rect width="200" height="60" fill="#eee"/><text x="10" y="35">수업 흐름도</text></svg>')

    # ── 사본 (내용이 똑같다)
    os.makedirs(os.path.join(root, "사본"), exist_ok=True)
    shutil.copy(os.path.join(root, "성적표.xlsx"),
                os.path.join(root, "사본", "성적표 - 복사본.xlsx"))

    # ── 압축 (한국 윈도우식 cp949 이름)
    make_cp949_zip(os.path.join(root, "압축묶음.zip"))

    # ── 윈도우 잡동사니 — 변환 대상이 아니다
    _write(os.path.join(root, "desktop.ini"), "[.ShellClassInfo]\nIconResource=x\n")
    _write(os.path.join(root, "Thumbs.db"), b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600, "wb")
    _write(os.path.join(root, "바로가기.lnk"), b"L\0\0\0\x01\x14\x02\0" + b"\0" * 80, "wb")

    # ── 아주 깊은 폴더 — 결과 경로가 260자를 넘는다
    deep = root
    for i in range(6):
        deep = os.path.join(deep, f"아주_긴_폴더_이름_{i}_" + "가" * 20)
    # 긴 경로를 지원하지 않는 PC 에서는 만들 수조차 없으므로 \\?\ 로 우회해 만든다.
    # 글자파일만이 아니라 PDF·워드도 둔다 — 각 라이브러리가 긴 경로를 여는지 봐야 한다.
    deep_long = "\\\\?\\" + os.path.abspath(deep) if os.name == "nt" else deep
    os.makedirs(deep_long, exist_ok=True)
    with io.open(os.path.join(deep_long, "깊은곳의_메모.txt"), "w", encoding="utf-8") as f:
        f.write("깊은 곳에 있는 메모입니다.\n")
    tmp_docx = os.path.join(tmp, "깊은곳의_워드.docx")
    make_docx(tmp_docx, img, "깊은 곳의 워드")
    shutil.copy(tmp_docx, os.path.join(deep_long, "깊은곳의_워드.docx"))
    import pymupdf
    tmp_pdf = os.path.join(tmp, "깊은곳의_PDF.pdf")
    d = pymupdf.open()
    d.new_page().insert_text((72, 90), "깊은 곳에 있는 PDF 본문", fontname="korea")
    d.save(tmp_pdf)
    d.close()
    shutil.copy(tmp_pdf, os.path.join(deep_long, "깊은곳의_PDF.pdf"))

    # ── 개인정보 함정
    _write(os.path.join(root, "개인정보", "함정.txt"),
           "프로젝트 이름 설정 방법을 알아봅니다. 환경 설정에서 바꿉니다.\n"
           "담당자 홍길동 (010-1234-5678)\n"
           "결제 안내 https://pay.example.com 카드번호 4111-1111-1111-1111\n"
           "계좌: 신한 110-123-456789\n"
           "계좌번호(경남은행) 207-0012-3456-78\n"
           "참고: doi: 10.1177/1365480216659733\n")
    _write(os.path.join(root, "개인정보", "홍길동 상담기록.txt"),
           "담당: 홍길동\n상담 내용입니다.\n")
    # 사람 이름이 파일 이름에 있고 그림도 든 문서 — 가리기가 그림 링크를 깨면 안 된다
    make_docx(os.path.join(root, "개인정보", "홍길동 사진자료.docx"), img, "담당자 홍길동")

    # ── 미리 만들어 둔 옛 형식·암호 오피스
    if os.path.isdir(STATIC):
        for dp, _, fns in os.walk(STATIC):
            for fn in fns:
                src = os.path.join(dp, fn)
                rel = os.path.relpath(src, STATIC)
                dst = os.path.join(root, "옛형식", rel)
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                shutil.copy(src, dst)

    shutil.rmtree(tmp)
    return root


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "_시험자료")
    build(out)
    n = sum(len(f) for _, _, f in os.walk(out))
    print(f"시험 자료 {n}개를 만들었습니다: {out}")
