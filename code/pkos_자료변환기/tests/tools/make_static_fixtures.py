# -*- coding: utf-8 -*-
"""
코드로 매번 만들 수 없는 시험 자료를 한 번만 만들어 tests/fixtures/ 에 둔다
===========================================================================

    옛엑셀.xls        xlwt 로 만든 진짜 BIFF8 파일 (정수·실수·날짜·빈 칸·오류 칸)
    암호워드.docx     msoffcrypto 로 암호(1234)를 건 워드 — 속은 OLE 의 EncryptedPackage

이 둘은 변환기가 아니라 **시험 자료를 만들 때만** 필요한 라이브러리를 쓴다.
평소 환경을 어지럽히지 않도록 임시 폴더에 깔아 쓰면 된다.

    pip install --target <임시폴더> xlwt msoffcrypto-tool
    set PYTHONPATH=<임시폴더>
    python tests/tools/make_static_fixtures.py

.doc / .ppt / .hwp 는 순수 파이썬으로 만들 방법이 없다. MS Office 가 저장 대화상자
없이 도는 PC 라면 make_office_fixtures.ps1 로 만들 수 있다. 없으면 자가진단은
그 항목을 '건너뜀'으로 표시한다.

내용은 모두 지어낸 것이다.
"""

import io
import os
import sys
import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "fixtures")
sys.path.insert(0, os.path.dirname(os.path.dirname(HERE)))


def make_xls(path: str):
    import xlwt
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("성적")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for c, h in enumerate(["번호", "과목", "점수", "날짜", "비고"]):
        ws.write(0, c, h)
    ws.write(1, 0, 1)
    ws.write(1, 1, "국어")
    ws.write(1, 2, 3500000.0)                 # 정수인데 실수로 저장되는 값
    ws.write(1, 3, datetime.date(2026, 3, 2), date_style)
    ws.write(2, 0, 2)
    ws.write(2, 1, "수학")
    ws.write(2, 2, 87.5)
    ws.write(2, 3, datetime.date(2026, 3, 3), date_style)
    ws.write(2, 4, xlwt.Formula("1/0"))       # #DIV/0! 오류 칸 (캐시 값 없음)
    hidden = wb.add_sheet("숨긴시트")
    hidden.write(0, 0, "숨긴 시트의 글")
    hidden.visibility = 1
    wb.save(path)


def make_encrypted_docx(path: str):
    import docx
    import msoffcrypto
    from msoffcrypto.format.ooxml import OOXMLFile
    d = docx.Document()
    d.add_paragraph("암호로 잠긴 문서의 본문")
    plain = io.BytesIO()
    d.save(plain)
    plain.seek(0)
    with open(path, "wb") as out:
        OOXMLFile(plain).encrypt("1234", out)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    make_xls(os.path.join(OUT, "옛엑셀.xls"))
    make_encrypted_docx(os.path.join(OUT, "암호워드.docx"))
    for fn in sorted(os.listdir(OUT)):
        print(f"  {fn:20} {os.path.getsize(os.path.join(OUT, fn)):>8,} 바이트")
