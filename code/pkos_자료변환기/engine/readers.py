# -*- coding: utf-8 -*-
"""
읽기 — 판별된 '형식'으로 알맞은 방법을 고른다
===============================================

sniff.py 가 **내용을 보고 정한 형식**을 받아 읽는다. 확장자는 보지 않는다.

처음 판은 한글·워드·PPT·엑셀·PDF·웹문서를 기존 엔진(readers_legacy.py)에 그대로
맡겼다. 일부러 까다로운 시험 자료를 만들어 돌려 보니 이런 것들이 새고 있었다.

    워드   표가 본문 끝으로 밀려남 · 글상자 글이 사라짐 · 서식파일(.dotx) 거부
    PPT    그룹으로 묶은 도형 안의 글이 사라짐 · 쇼 파일(.ppsx) 거부
    엑셀   이름이 .xlsx 가 아니면 openpyxl 이 확장자만 보고 거부
    웹     <div> 로만 된 글이 사라짐 · 표가 칸마다 한 줄로 흩어짐 · <pre> 코드 사라짐
    한글   서식이 바뀌는 곳마다 낱말 중간에 공백이 끼어듦 · 표 속 표가 바깥 칸을 덮어씀
    PDF    암호 걸린 PDF 를 '스캔 문서'라고 안내
    탭 구분 표(TSV) 를 쉼표로 쪼개 한 칸짜리 표로 만듦

그래서 이 모듈이 직접 읽는다. 기존 엔진에서는 **검증이 끝난 부분만** 가져다 쓴다
— 한글 .hwp 이진 해석(표 복원 포함)이 그렇다.

그림은 **만난 자리에** 적는다. 읽기 함수가 그림을 만나면 ReadContext.save_image()
에 넘기고 그 자리에 링크를 둔다. (예전에는 모두 문서 맨 끝에 모았다)

PKOS(개인지식운영체계) 프로젝트
"""

from __future__ import annotations

import io
import os
import re
import csv
import json
import zlib
import struct
import hashlib
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field

from .readers_legacy import ReadResult, read_hwp as _legacy_read_hwp
from .oldoffice import read_doc, read_ppt, read_xls
from .fsutil import long_path, md_link, write_bytes, copy_file
from .mdutil import rows_to_md, fmt_value, join_blocks
from . import images as _img
from . import ocr as _ocr


# ═════════════════════════════════════════════════════════════
# 그림 저장 자리
# ═════════════════════════════════════════════════════════════
@dataclass
class ReadContext:
    """읽는 동안 만난 그림을 어디에 어떻게 저장할지."""
    asset_dir: str = ""          # 그림을 저장할 폴더. 비우면 그림을 꺼내지 않는다
    asset_url: str = ""          # .md 에서 그 폴더를 가리키는 상대 경로
    min_bytes: int = 8000        # 이보다 작으면 아이콘·글머리표로 본다
    min_px: int = 48             # 가로나 세로가 이보다 작으면 버린다
    limit: int = 300
    saved: list = field(default_factory=list)
    _seen: dict = field(default_factory=dict)

    @property
    def enabled(self) -> bool:
        return bool(self.asset_dir)

    def save_image(self, data: bytes, ext: str = "", alt: str = "") -> str:
        """그림을 저장하고 그 자리에 적을 마크다운 한 줄을 돌려준다."""
        if not self.enabled or not data or len(data) < self.min_bytes:
            return ""
        key = hashlib.sha1(data).hexdigest()
        if key in self._seen:                        # 같은 그림을 여러 번 쓴 문서
            fn = self._seen[key]
        else:
            if len(self.saved) >= self.limit:
                return ""
            data2, ext2 = _img.web_ready(data, ext)
            if data2 is None:
                return ""
            wh = _img.image_size(data2)
            if wh and min(wh) < self.min_px:
                return ""
            fn = f"{len(self.saved) + 1:03d}.{ext2}"
            write_bytes(os.path.join(self.asset_dir, fn), data2)
            self.saved.append(fn)
            self._seen[key] = fn
        return md_link(alt or "그림", f"{self.asset_url}/{fn}", image=True)

    def save_original(self, path: str, name: str) -> str:
        """사진 파일 자체를 그대로 옮긴다 (크기·형식을 따지지 않는다)."""
        if not self.enabled:
            return ""
        copy_file(path, os.path.join(self.asset_dir, name))
        self.saved.append(name)
        return f"{self.asset_url}/{name}"


# ═════════════════════════════════════════════════════════════
# 글자파일
# ═════════════════════════════════════════════════════════════
def _load_text(path: str, encoding: str = "") -> tuple[str, str]:
    """(본문, 실제로 쓴 인코딩)."""
    with open(long_path(path), "rb") as f:
        raw = f.read()
    tries = [encoding] if encoding else []
    tries += ["utf-8-sig", "cp949", "utf-16"]
    for enc in tries:
        try:
            return raw.decode(enc), enc
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", "replace"), "utf-8(일부 깨짐)"


def read_plain(path: str, encoding: str = "") -> ReadResult:
    text, enc = _load_text(path, encoding)
    text = text.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    return ReadResult(True, text.strip(), "텍스트", {"인코딩": enc})


_FRONT = re.compile(r"\A---[ \t]*\n(.*?)\n(?:---|\.\.\.)[ \t]*(?:\n|\Z)", re.S)


def _parse_tags(front: str) -> list[str]:
    """머리말에서 tags 만 꺼낸다 — tags: [a, b] 와 목록형 둘 다."""
    m = re.search(r"^tags\s*:\s*\[(.*?)\]\s*$", front, re.M)
    if m:
        return [t.strip().strip("'\"") for t in m.group(1).split(",") if t.strip()]
    m = re.search(r"^tags\s*:\s*\n((?:[ \t]*-[^\n]*\n?)+)", front, re.M)
    if m:
        return [re.sub(r"^[ \t]*-\s*", "", l).strip().strip("'\"")
                for l in m.group(1).splitlines() if l.strip()]
    m = re.search(r"^tags\s*:\s*(\S.*)$", front, re.M)
    if m:
        return [t.strip().lstrip("#") for t in re.split(r"[,\s]+", m.group(1)) if t.strip()]
    return []


def read_markdown(path: str, encoding: str = "") -> ReadResult:
    """마크다운. 원래 머리말(YAML)이 있으면 본문에서 떼어 따로 넘긴다
    (그대로 두면 새 머리말 뒤에 --- 가 또 나와 본문이 깨져 보인다)."""
    r = read_plain(path, encoding)
    r.kind = "마크다운"
    m = _FRONT.match(r.text)
    if m:
        r.meta["원래머리말"] = m.group(1)
        tags = _parse_tags(m.group(1))
        if tags:
            r.meta["tags"] = tags
        tm = re.search(r"^title\s*:\s*(.+)$", m.group(1), re.M)
        if tm:
            r.meta["제목후보"] = tm.group(1).strip().strip("'\"")
        r.text = r.text[m.end():].strip()
    return r


def read_json(path: str, encoding: str = "") -> ReadResult:
    text, enc = _load_text(path, encoding)
    try:
        obj = json.loads(text.lstrip("﻿"))
    except Exception as e:
        return ReadResult(True, f"```\n{text[:200000]}\n```", "JSON",
                          {"인코딩": enc, "해석": f"실패: {e}"})
    if isinstance(obj, dict) and "cells" in obj and "nbformat" in obj:
        lang = (obj.get("metadata", {}).get("kernelspec", {}) or {}).get("language", "python")
        blocks = []
        for c in obj.get("cells", []):
            src = c.get("source", "")
            src = "".join(src) if isinstance(src, list) else str(src)
            if not src.strip():
                continue
            blocks.append(src.strip() if c.get("cell_type") == "markdown"
                          else f"```{lang}\n{src.rstrip()}\n```")
        return ReadResult(True, join_blocks(blocks), "주피터 노트북",
                          {"셀": len(obj.get("cells", []))})
    pretty = json.dumps(obj, ensure_ascii=False, indent=2)
    if len(pretty) > 200000:
        pretty = pretty[:200000] + "\n… (너무 길어 줄임)"
    return ReadResult(True, f"```json\n{pretty}\n```", "JSON", {"인코딩": enc})


def read_srt(path: str, encoding: str = "") -> ReadResult:
    """자막에서 번호·시간표·꾸밈 태그를 걷어내고 대사만 남긴다 (.srt .vtt .smi)."""
    text, enc = _load_text(path, encoding)
    lines, n = [], 0
    for ln in text.splitlines():
        s = ln.strip()
        if not s or s.isdigit() or s.upper().startswith(("WEBVTT", "NOTE")):
            continue
        if "-->" in s or re.match(r"(?i)<sync\b", s):
            n += 1
        s = re.sub(r"<[^>]{1,60}>", "", s)
        s = s.replace("&nbsp;", " ").strip()
        if s and "-->" not in s:
            lines.append(s)
    body = re.sub(r"\s{2,}", " ", " ".join(lines)).strip()
    return ReadResult(True, body, "자막", {"자막수": n, "인코딩": enc})


_CODE_LANG = {".py": "python", ".js": "javascript", ".ts": "typescript", ".java": "java",
              ".c": "c", ".cpp": "cpp", ".h": "c", ".cs": "csharp", ".r": "r", ".sql": "sql",
              ".sh": "bash", ".css": "css", ".yml": "yaml", ".yaml": "yaml", ".toml": "toml",
              ".bat": "bat", ".cmd": "bat", ".ps1": "powershell"}


def read_code(path: str, encoding: str = "") -> ReadResult:
    text, enc = _load_text(path, encoding)
    lang = _CODE_LANG.get(os.path.splitext(path)[1].lower(), "")
    fence = "````" if "```" in text else "```"
    return ReadResult(True, f"{fence}{lang}\n{text.rstrip()}\n{fence}", "소스코드",
                      {"줄": text.count("\n") + 1, "인코딩": enc})


def read_delimited(path: str, encoding: str = "", delimiter: str = ",") -> ReadResult:
    text, enc = _load_text(path, encoding)
    rows = list(csv.reader(io.StringIO(text.lstrip("﻿")), delimiter=delimiter))
    return ReadResult(True, rows_to_md(rows), "표 데이터", {"행": len(rows), "인코딩": enc})


# ── 웹문서 ───────────────────────────────────────────────────
_SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "iframe", "object",
              "embed", "head", "meta", "link", "button", "input", "select", "textarea",
              "canvas", "map"}
_BLOCK_TAGS = {"p", "div", "section", "article", "main", "header", "footer", "aside",
               "nav", "figure", "figcaption", "address", "form", "fieldset", "details",
               "summary", "dl", "dt", "dd", "center", "body", "html", "li", "caption"}


def _html_encoding(raw: bytes, fallback: str) -> str:
    m = re.search(rb"<meta[^>]+charset\s*=\s*[\"']?\s*([A-Za-z0-9_\-]+)", raw[:8192], re.I)
    if m:
        enc = m.group(1).decode("ascii", "ignore").lower()
        enc = {"euc-kr": "cp949", "ks_c_5601-1987": "cp949", "x-windows-949": "cp949",
               "ms949": "cp949"}.get(enc, enc)
        try:
            "".encode(enc)
            return enc
        except LookupError:
            pass
    return fallback or "utf-8"


def html_to_markdown(html: str) -> str:
    from bs4 import BeautifulSoup, NavigableString, Comment, Tag
    soup = BeautifulSoup(html, "html.parser")
    out: list[str] = []
    buf: list[str] = []

    def flush():
        s = re.sub(r"[ \t\r\f\v]+", " ", "".join(buf))
        s = "\n".join(x.strip() for x in s.split("\n")).strip()
        if s:
            out.append(s)
        buf.clear()

    def text_of(node) -> str:
        return re.sub(r"\s+", " ", node.get_text(" ", strip=True)).strip()

    def lists(node, depth: int, lines: list[str]):
        for li in node.find_all("li", recursive=False):
            parts = []
            for ch in li.children:
                if isinstance(ch, Tag) and ch.name in ("ul", "ol"):
                    continue
                parts.append(text_of(ch) if isinstance(ch, Tag) else str(ch))
            t = re.sub(r"\s+", " ", " ".join(parts)).strip()
            if t:
                lines.append("  " * depth + "- " + t)
            for sub in li.find_all(["ul", "ol"], recursive=False):
                lists(sub, depth + 1, lines)

    def table(tb) -> str:
        rows = []
        for tr in tb.find_all("tr"):
            if tr.find_parent("table") is not tb:
                continue
            cells = []
            for td in tr.find_all(["td", "th"], recursive=False):
                cells.append(text_of(td))
                span = str(td.get("colspan", "1"))
                cells += [""] * (min(int(span), 30) - 1 if span.isdigit() else 0)
            if cells:
                rows.append(cells)
        return rows_to_md(rows)

    def walk(node):
        for ch in node.children:
            if isinstance(ch, Comment):
                continue
            if isinstance(ch, NavigableString):
                buf.append(str(ch))
                continue
            if not isinstance(ch, Tag):
                continue
            name = (ch.name or "").lower()
            if name in _SKIP_TAGS:
                continue
            if name == "br":
                buf.append("\n")
            elif re.fullmatch(r"h[1-6]", name):
                flush()
                t = text_of(ch)
                if t:
                    out.append("#" * int(name[1]) + " " + t)
            elif name in ("ul", "ol"):
                flush()
                lines: list[str] = []
                lists(ch, 0, lines)
                if lines:
                    out.append("\n".join(lines))
            elif name == "table":
                flush()
                t = table(ch)
                if t:
                    out.append(t)
            elif name == "pre":
                flush()
                code = ch.get_text()
                if code.strip():
                    out.append("```\n" + code.strip("\n") + "\n```")
            elif name == "blockquote":
                flush()
                inner = html_to_markdown(str(ch.decode_contents()))
                if inner:
                    out.append("\n".join("> " + l if l else ">" for l in inner.splitlines()))
            elif name == "hr":
                flush()
            elif name == "img":
                alt = (ch.get("alt") or "").strip()
                if alt:
                    buf.append(f" {alt} ")
            elif name in _BLOCK_TAGS:
                flush()
                walk(ch)
                flush()
            else:
                walk(ch)

    walk(soup)
    flush()
    return join_blocks(out)


def read_html(path: str, encoding: str = "") -> ReadResult:
    with open(long_path(path), "rb") as f:
        raw = f.read()
    enc = _html_encoding(raw, encoding)
    html = raw.decode(enc, "replace")
    title = ""
    m = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    if m:
        title = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", m.group(1))).strip()
    try:
        body = html_to_markdown(html)
    except ImportError:
        body = re.sub(r"\s{2,}", " ", re.sub(r"<[^>]+>", " ", html)).strip()
    meta = {"인코딩": enc}
    if title:
        meta["제목후보"] = title
    return ReadResult(True, body, "웹문서", meta)


def read_xml(path: str, encoding: str = "") -> ReadResult:
    text, enc = _load_text(path, encoding)
    try:
        root = ET.fromstring(text.lstrip("﻿").encode("utf-8"))
        parts = [t.strip() for t in root.itertext() if t and t.strip()]
        body = "\n".join(parts)
    except Exception:
        body = re.sub(r"\s{2,}", " ", re.sub(r"<[^>]+>", " ", text)).strip()
    if len(body) < 3:
        return ReadResult(False, kind="XML", error="글자를 찾지 못했습니다")
    return ReadResult(True, body, "XML", {"인코딩": enc})


# ═════════════════════════════════════════════════════════════
# 오피스 (새 형식) — 서식파일·쇼파일·매크로 파일도 연다
# ═════════════════════════════════════════════════════════════
_CT_MAIN = {
    "docx": b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "pptx": b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
}
_CT_VARIANTS = {
    "docx": [b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
             b"application/vnd.ms-word.document.macroEnabled.main+xml",
             b"application/vnd.ms-word.template.macroEnabledTemplate.main+xml"],
    "pptx": [b"application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
             b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
             b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
             b"application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml",
             b"application/vnd.ms-powerpoint.template.macroEnabled.main+xml"],
}


def _open_ooxml(path: str, kind: str):
    """python-docx/pptx 는 '문서 본체 종류'가 정확히 문서여야만 연다.
    서식파일(.dotx)·쇼(.ppsx)·매크로(.pptm) 는 내용이 같은데도 거부하므로,
    메모리 안에서 종류 표시만 문서로 바꿔 연다. 원본 파일은 건드리지 않는다."""
    opener = __import__("docx").Document if kind == "docx" else __import__("pptx").Presentation
    with open(long_path(path), "rb") as fh:
        data = fh.read()
    try:
        return opener(io.BytesIO(data))
    except Exception:
        fixed = io.BytesIO()
        changed = False
        with zipfile.ZipFile(io.BytesIO(data)) as zin, \
                zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                blob = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    for v in _CT_VARIANTS[kind]:
                        if v in blob:
                            blob = blob.replace(v, _CT_MAIN[kind])
                            changed = True
                zout.writestr(item, blob)
        if not changed:
            raise
        fixed.seek(0)
        return opener(fixed)


# ── 워드 ─────────────────────────────────────────────────────
_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def _q(tag: str) -> str:
    p, local = tag.split(":")
    return f"{{{_NS[p]}}}{local}"


def _in_fallback(el, stop) -> bool:
    """mc:Fallback 안의 요소인가 — 새 모양(Choice)과 옛 모양(Fallback)에 같은
    글상자·그림이 두 번 들어 있으므로 하나만 읽어야 한다."""
    p = el.getparent()
    while p is not None and p is not stop:
        if p.tag == _q("mc:Fallback"):
            return True
        p = p.getparent()
    return False


def _docx_heading_level(para, p_el) -> int:
    ppr = p_el.find(_q("w:pPr"))
    if ppr is not None:
        ol = ppr.find(_q("w:outlineLvl"))
        if ol is not None and str(ol.get(_q("w:val"), "")).isdigit():
            lvl = int(ol.get(_q("w:val"))) + 1
            if lvl <= 6:
                return lvl
    try:
        name = (para.style.name or "").lower() if para.style is not None else ""
    except Exception:
        name = ""
    if name == "title":
        return 1
    m = re.match(r"(?:heading|제목)\s*(\d)", name)
    return min(int(m.group(1)), 6) if m else 0


def _docx_paragraph(p_el, doc, ctx: ReadContext) -> list[str]:
    from docx.text.paragraph import Paragraph
    para = Paragraph(p_el, doc)
    out = []
    text = (para.text or "").strip()
    if text:
        lvl = _docx_heading_level(para, p_el)
        ppr = p_el.find(_q("w:pPr"))
        num = ppr.find(_q("w:numPr")) if ppr is not None else None
        if lvl:
            out.append("#" * lvl + " " + text.replace("\n", " "))
        elif num is not None:
            ilvl = num.find(_q("w:ilvl"))
            depth = int(ilvl.get(_q("w:val"), "0")) if ilvl is not None else 0
            out.append("  " * depth + "- " + text)
        else:
            out.append(text)

    # 글상자 (본문 흐름 밖에 떠 있는 글)
    for tb in p_el.iter(_q("w:txbxContent")):
        if _in_fallback(tb, p_el):
            continue
        lines = ["".join(t.text or "" for t in pp.iter(_q("w:t"))).strip()
                 for pp in tb.iter(_q("w:p"))]
        lines = [l for l in lines if l]
        if lines:
            out.append("\n".join(lines))

    # 그림 (새 모양 a:blip, 옛 모양 v:imagedata)
    if ctx.enabled:
        rels = doc.part.related_parts
        for el, attr in [(b, _q("r:embed")) for b in p_el.iter(_q("a:blip"))] + \
                        [(v, _q("r:id")) for v in p_el.iter(_q("v:imagedata"))]:
            if _in_fallback(el, p_el):
                continue
            part = rels.get(el.get(attr))
            if part is None:
                continue
            alt = ""
            for dp in p_el.iter(_q("wp:docPr")):
                alt = dp.get("descr") or dp.get("title") or ""
                break
            line = ctx.save_image(part.blob, os.path.splitext(str(part.partname))[1], alt)
            if line:
                out.append(line)
    return out


def _docx_table(tbl_el, doc) -> str:
    from docx.table import Table
    rows = []
    for r in Table(tbl_el, doc).rows:
        cells, prev = [], None
        for c in r.cells:
            if c._tc is prev:                       # 가로로 합친 칸은 한 번만
                cells.append("")
                continue
            prev = c._tc
            parts = [p.text for p in c.paragraphs]
            for inner in c._tc.iterchildren(_q("w:tbl")):   # 칸 안의 표
                parts.append(" / ".join(
                    "".join(t.text or "" for t in tc.iter(_q("w:t"))).strip()
                    for tc in inner.iter(_q("w:tc"))))
            cells.append("\n".join(x for x in parts if x.strip()))
        rows.append(cells)
    return rows_to_md(rows)


def _docx_notes(path: str) -> list[str]:
    blocks = []
    W = _NS["w"]
    try:
        with zipfile.ZipFile(long_path(path)) as z:
            names = set(z.namelist())
            for name, label, tag in (("word/footnotes.xml", "각주", "footnote"),
                                     ("word/endnotes.xml", "미주", "endnote")):
                if name not in names:
                    continue
                root = ET.fromstring(z.read(name))
                items = []
                for fn in root.findall(f"{{{W}}}{tag}"):
                    if fn.get(f"{{{W}}}type") in ("separator", "continuationSeparator",
                                                  "continuationNotice"):
                        continue
                    t = " ".join("".join(x.text or "" for x in p.iter(f"{{{W}}}t")).strip()
                                 for p in fn.iter(f"{{{W}}}p")).strip()
                    if t:
                        items.append(f"{len(items) + 1}. {t}")
                if items:
                    blocks.append(f"## {label}\n\n" + "\n".join(items))
    except Exception:
        pass
    return blocks


def read_docx(path: str, ctx: ReadContext) -> ReadResult:
    doc = _open_ooxml(path, "docx")
    blocks: list[str] = []
    n_tables = 0

    def handle(el):
        nonlocal n_tables
        if el.tag == _q("w:p"):
            blocks.extend(_docx_paragraph(el, doc, ctx))
        elif el.tag == _q("w:tbl"):
            t = _docx_table(el, doc)
            if t:
                blocks.append(t)
                n_tables += 1
        elif el.tag == _q("w:sdt"):                  # 내용 컨트롤 — 안에 문단·표가 있다
            content = el.find(_q("w:sdtContent"))
            if content is not None:
                for sub in content.iterchildren():
                    handle(sub)

    for child in doc.element.body.iterchildren():
        handle(child)

    body_text = "\n".join(blocks)
    extra = []
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            try:
                if hf.is_linked_to_previous:
                    continue
                for p in hf.paragraphs:
                    t = p.text.strip()
                    if t and t not in extra and t not in body_text:
                        extra.append(t)
            except Exception:
                continue
    if extra:
        blocks.append("## 머리글·바닥글\n\n" + "\n".join(extra))
    blocks += _docx_notes(path)
    return ReadResult(True, join_blocks(blocks), "워드",
                      {"표": n_tables, "그림": len(ctx.saved)})


# ── 파워포인트 ───────────────────────────────────────────────
def _pptx_shape(sh, out: list[str], ctx: ReadContext, depth: int = 0):
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    try:
        st = sh.shape_type
    except Exception:
        st = None
    if st == MSO_SHAPE_TYPE.GROUP and depth < 10:
        for c in sh.shapes:
            _pptx_shape(c, out, ctx, depth + 1)
        return
    if getattr(sh, "has_table", False) and sh.has_table:
        rows = [[c.text for c in r.cells] for r in sh.table.rows]
        t = rows_to_md(rows)
        if t:
            out.append(t)
        return
    if getattr(sh, "has_chart", False) and sh.has_chart:
        out.extend(_pptx_chart(sh.chart))
        return
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        bullets = False
        try:
            bullets = sh.is_placeholder and sh.placeholder_format.type in (
                PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
        except Exception:
            pass
        lines = []
        for para in sh.text_frame.paragraphs:
            t = re.sub(r"[\v\r\n]+", " ", para.text or "").strip()
            if t:
                lines.append(("  " * para.level + "- " + t) if bullets else t)
        if lines:
            out.append("\n".join(lines))
    if ctx.enabled:
        try:
            img = sh.image                           # 그림, 그림이 채워진 자리표시
        except Exception:
            img = None
        if img is not None:
            alt = ""
            try:
                alt = sh._element.xpath("./*[1]/p:cNvPr/@descr")[0]
            except Exception:
                pass
            line = ctx.save_image(img.blob, img.ext, alt)
            if line:
                out.append(line)


def _pptx_chart(chart) -> list[str]:
    out = []
    try:
        if chart.has_title:
            t = chart.chart_title.text_frame.text.strip()
            if t:
                out.append(f"**차트: {t}**")
    except Exception:
        pass
    try:
        plot = chart.plots[0]
        cats = list(plot.categories)
        series = list(plot.series)
        rows = [["항목"] + [s.name for s in series]]
        for i, c in enumerate(cats):
            rows.append([fmt_value(c)] + [fmt_value(list(s.values)[i]) if i < len(list(s.values))
                                          else "" for s in series])
        t = rows_to_md(rows)
        if t:
            out.append(t)
    except Exception:
        pass
    return out


def read_pptx(path: str, ctx: ReadContext) -> ReadResult:
    prs = _open_ooxml(path, "pptx")
    blocks, n_notes = [], 0
    for n, slide in enumerate(prs.slides, 1):
        title_sh = None
        try:
            title_sh = slide.shapes.title
        except Exception:
            pass
        title = ""
        if title_sh is not None and title_sh.has_text_frame:
            title = re.sub(r"\s+", " ", title_sh.text_frame.text).strip()
        blocks.append(f"## {n}. {title}" if title else f"## 슬라이드 {n}")
        for sh in slide.shapes:
            if title_sh is not None and sh.shape_id == title_sh.shape_id:
                continue
            _pptx_shape(sh, blocks, ctx)
        try:
            if slide.has_notes_slide:
                tf = slide.notes_slide.notes_text_frame
                note = tf.text.strip() if tf is not None else ""
                if note:
                    n_notes += 1
                    blocks.append("> **발표자 노트**\n" +
                                  "\n".join("> " + l for l in note.splitlines() if l.strip()))
        except Exception:
            pass
    return ReadResult(True, join_blocks(blocks), "파워포인트",
                      {"슬라이드": len(prs.slides), "노트": n_notes, "그림": len(ctx.saved)})


# ── 엑셀 ─────────────────────────────────────────────────────
def read_xlsx(path: str, ctx: ReadContext) -> ReadResult:
    import openpyxl
    blocks, n_rows = [], 0
    # 경로 대신 파일 객체를 넘긴다 — openpyxl 은 경로를 받으면 확장자부터 검사해서
    # 이름이 .xlsx 가 아닌 엑셀 파일을 거부한다.
    with open(long_path(path), "rb") as fh:
        wb = openpyxl.load_workbook(fh, data_only=True, read_only=True)
        try:
            for ws in wb.worksheets:
                rows = [[fmt_value(v) for v in row] for row in ws.iter_rows(values_only=True)]
                n_rows += len(rows)
                t = rows_to_md(rows)
                if t:
                    hidden = " (숨긴 시트)" if getattr(ws, "sheet_state", "visible") != "visible" else ""
                    blocks += [f"## {ws.title}{hidden}", t]
            n_sheets = len(wb.worksheets)
        finally:
            wb.close()
    if ctx.enabled:
        shots = [ctx.save_image(data, os.path.splitext(name)[1], name)
                 for name, data in _img.zip_images(path, ("xl/media/",))]
        shots = [s for s in shots if s]
        if shots:
            blocks.append("## 시트 속 그림\n\n" + "\n\n".join(shots))
    return ReadResult(True, join_blocks(blocks), "엑셀", {"시트": n_sheets, "행": n_rows})


# ═════════════════════════════════════════════════════════════
# PDF
# ═════════════════════════════════════════════════════════════
def _pymupdf():
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        import fitz as pymupdf
        return pymupdf


def _pdf_page_text(page) -> str:
    """쪽의 글 덩어리를 문단으로. 줄 끝마다 끊긴 것을 이어 붙인다."""
    paras = []
    for b in page.get_text("blocks"):
        if len(b) < 7 or b[6] != 0:                 # 6번 값 0 = 글, 1 = 그림
            continue
        t = b[4].replace("-\n", "").replace("\n", " ")
        t = re.sub(r"\s{2,}", " ", t).strip()
        if t:
            paras.append(t)
    return "\n\n".join(paras)


def read_pdf(path: str, ctx: ReadContext) -> ReadResult:
    pymupdf = _pymupdf()
    kind = "PDF"
    try:
        doc = pymupdf.open(long_path(path), filetype="pdf")
    except Exception as e:
        return ReadResult(False, kind=kind, error=f"PDF 를 열지 못했습니다: {e}")
    try:
        if doc.needs_pass:
            return ReadResult(False, kind=kind,
                              error="암호가 걸린 PDF 입니다. 암호를 풀어 저장한 뒤 넣어 주세요.")
        pages = doc.page_count
        blocks, empty, total = [], [], 0
        for pno in range(pages):
            page = doc[pno]
            text = _pdf_page_text(page)
            total += len(text)
            page_blocks = [f"<!-- 쪽 {pno + 1} -->"]
            if text:
                page_blocks.append(text)
            else:
                empty.append(pno + 1)
            if ctx.enabled:
                for info in page.get_images(full=True):
                    try:
                        img = doc.extract_image(info[0])
                        data, ext = img.get("image", b""), img.get("ext", "")
                        if ext.lower() not in _img.WEB_EXT:
                            pix = pymupdf.Pixmap(doc, info[0])
                            if pix.n - pix.alpha >= 4:
                                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                            data, ext = pix.tobytes("png"), "png"
                    except Exception:
                        continue
                    line = ctx.save_image(data, ext, f"{pno + 1}쪽 그림")
                    if line:
                        page_blocks.append(line)
            blocks.append("\n\n".join(page_blocks))

        if total < 20 and pages:
            # 글자가 거의 없을 때 — 그림이 있으면 스캔본이고, 그림도 없으면 짧은 글뿐인
            # 정상 PDF(한 줄짜리 증명서 따위)다. 처음 판은 글자 수만 보고 둘 다 거절했다.
            has_images = any(doc[p].get_images() for p in range(min(pages, 20)))
            if has_images:
                if _ocr.is_on():
                    return _ocr_pdf(doc, ctx, pymupdf)
                return ReadResult(False, kind=kind,
                                  error=(f"글자가 없는 PDF 입니다({pages}쪽). 스캔하거나 사진으로 "
                                         f"만든 문서로 보입니다. --글자인식 을 켜면 읽을 수 있습니다."))
            if total == 0:
                return ReadResult(False, kind=kind, error=f"빈 PDF 입니다({pages}쪽, 글자도 그림도 없음).")
        meta = {"쪽": pages}
        if empty:
            meta["글자없는쪽"] = len(empty)
        return ReadResult(True, join_blocks(blocks), kind, meta)
    finally:
        doc.close()


def _ocr_pdf(doc, ctx: ReadContext, pymupdf) -> ReadResult:
    """스캔 PDF — 쪽마다 그림으로 그려 글자 인식을 돌린다."""
    blocks, ok_pages, err = [], 0, ""
    for pno in range(doc.page_count):
        pix = doc[pno].get_pixmap(dpi=200)
        r = _ocr.run_ocr_bytes(pix.tobytes("png"), "png")
        blocks.append(f"<!-- 쪽 {pno + 1} -->")
        if r.ok and r.text.strip():
            blocks.append(r.text.strip())
            ok_pages += 1
        else:
            err = r.error
    if not ok_pages:
        return ReadResult(False, kind="PDF(스캔)",
                          error=f"스캔 PDF 에서 글자 인식에 실패했습니다: {err}")
    return ReadResult(True, join_blocks(blocks), "PDF(스캔)",
                      {"쪽": doc.page_count, "글자인식": _ocr.current_engine()})


# ═════════════════════════════════════════════════════════════
# 한글
# ═════════════════════════════════════════════════════════════
def read_hwp(path: str, ctx: ReadContext) -> ReadResult:
    """구형 .hwp — 암호·배포용 문서는 미리 알아보고 이유를 알린다.
    본문 해석은 검증된 기존 엔진을 쓰고, 그림은 문서 끝에 모은다
    (이진 한글에서 그림의 제자리를 찾으려면 개체 레코드를 모두 풀어야 한다)."""
    try:
        import olefile
        ole = olefile.OleFileIO(long_path(path))
        try:
            header = ole.openstream("FileHeader").read() if ole.exists("FileHeader") else b""
        finally:
            ole.close()
        if len(header) >= 40:
            props = struct.unpack_from("<I", header, 36)[0]
            if props & 0x02:
                return ReadResult(False, kind="한글",
                                  error="암호가 걸린 한글 문서입니다. 한글에서 암호를 풀고 "
                                        "저장한 뒤 넣어 주세요.")
            if props & 0x04:
                return ReadResult(False, kind="한글",
                                  error="배포용 한글 문서입니다. 내용이 암호화되어 있어 읽을 수 "
                                        "없습니다. 한글에서 열어 PDF 로 인쇄한 뒤 넣어 주세요.")
    except Exception:
        pass

    res = _legacy_read_hwp(long_path(path))
    if res.ok and ctx.enabled:
        shots = [ctx.save_image(data, os.path.splitext(name)[1], name)
                 for name, data in _img.hwp_images(path)]
        shots = [s for s in shots if s]
        if shots:
            res.text = join_blocks([res.text, "## 문서 속 그림\n\n" + "\n\n".join(shots)])
            res.meta["그림"] = len(shots)
    return res


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _hwpx_text(el) -> str:
    """<hp:t> 한 덩이의 글. 안에 든 탭·줄바꿈 요소 뒤의 글(tail)까지 챙긴다."""
    parts = [el.text or ""]
    for ch in el:
        name = _local(ch.tag)
        if name == "lineBreak":
            parts.append("\n")
        elif name == "tab":
            parts.append("\t")
        parts.append(ch.tail or "")
    return "".join(parts)


def read_hwpx(path: str, ctx: ReadContext) -> ReadResult:
    kind = "한글"
    try:
        z = zipfile.ZipFile(long_path(path))
    except Exception as e:
        return ReadResult(False, kind=kind, error=f"파일 열기 실패: {e}")
    try:
        names = z.namelist()
        secs = sorted((n for n in names if re.match(r"Contents/section\d+\.xml$", n)),
                      key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))
        if not secs:
            return ReadResult(False, kind=kind, error="본문(section XML)을 찾지 못했습니다")

        bin_map = {}                                  # 그림 id → BinData 경로
        if "Contents/content.hpf" in names:
            try:
                for it in ET.fromstring(z.read("Contents/content.hpf")).iter():
                    href = it.get("href", "")
                    if _local(it.tag) == "item" and "BinData/" in href:
                        bin_map[it.get("id")] = href if href in names else "Contents/" + href
            except Exception:
                pass

        blocks: list[str] = []
        n_tables = 0

        def cell_text(tc) -> str:
            parts: list[str] = []

            def walk_cell(el, buf):
                name = _local(el.tag)
                if name == "tbl":                     # 칸 안의 표 — 칸 글로 풀어 잇는다
                    inner = [cell_text(t) for t in el.iter() if _local(t.tag) == "tc" and t is not el]
                    buf.append(" / ".join(x for x in inner if x))
                    return
                if name == "t":
                    buf.append(_hwpx_text(el))
                    return
                for ch in el:
                    walk_cell(ch, buf)
                if name == "p":
                    s = "".join(buf).strip()
                    if s:
                        parts.append(s)
                    buf.clear()

            walk_cell(tc, [])
            return "\n".join(parts)

        def table_md(tbl) -> str:
            cells, spans = {}, {}
            for tr in (c for c in tbl if _local(c.tag) == "tr"):
                for tc in (c for c in tr if _local(c.tag) == "tc"):
                    addr = next((a for a in tc if _local(a.tag) == "cellAddr"), None)
                    if addr is None:
                        continue
                    try:
                        r, c = int(addr.get("rowAddr", 0)), int(addr.get("colAddr", 0))
                    except ValueError:
                        continue
                    if r > 500 or c > 200:
                        continue
                    cells[(r, c)] = cell_text(tc)
            if not cells:
                return ""
            nr = max(r for r, _ in cells) + 1
            nc = max(c for _, c in cells) + 1
            grid = [["" for _ in range(nc)] for _ in range(nr)]
            for (r, c), v in cells.items():
                grid[r][c] = v
            return rows_to_md(grid)

        def walk(el, buf):
            nonlocal n_tables
            name = _local(el.tag)
            if name == "tbl":
                s = "".join(buf).strip()
                if s:
                    blocks.append(s)
                buf.clear()
                t = table_md(el)
                if t:
                    blocks.append(t)
                    n_tables += 1
                return
            if name == "pic" and ctx.enabled:
                ref = next((x.get("binaryItemIDRef") for x in el.iter()
                            if _local(x.tag) == "img"), None)
                href = bin_map.get(ref)
                if href and href in names:
                    line = ctx.save_image(z.read(href), os.path.splitext(href)[1])
                    if line:
                        s = "".join(buf).strip()
                        if s:
                            blocks.append(s)
                        buf.clear()
                        blocks.append(line)
            if name == "t":
                buf.append(_hwpx_text(el))
                return
            for ch in el:
                walk(ch, buf)
            if name == "p":
                s = "".join(buf).strip()
                if s:
                    blocks.append(s)
                buf.clear()

        for s in secs:
            leftover: list[str] = []
            walk(ET.fromstring(z.read(s)), leftover)
            if "".join(leftover).strip():
                blocks.append("".join(leftover).strip())
        return ReadResult(True, join_blocks(blocks), kind,
                          {"섹션": len(secs), "표": n_tables, "그림": len(ctx.saved)})
    finally:
        z.close()


# ── 한글 HML (한글 문서의 XML 판, HWPML 2.x) ──────────────────
def read_hml(path: str, ctx: ReadContext) -> ReadResult:
    """<HWPML><BODY><SECTION><P><TEXT><CHAR>글</CHAR><TABLE>…</TABLE></TEXT></P>…

    그냥 XML 로 읽으면 머리(HEAD)의 글꼴·스타일 이름까지 본문에 섞인다.
    BODY 만 훑고, 표는 칸 주소(RowAddr·ColAddr)로 되살린다.
    그림은 꼬리(TAIL)의 BINDATA 에 base64 로 들어 있다.
    """
    import base64
    try:
        root = ET.parse(long_path(path)).getroot()
    except Exception as e:
        return ReadResult(False, kind="한글(HML)", error=f"HML 을 해석하지 못했습니다: {e}")

    items = [it for it in root.iter("BINITEM")]           # BinItem 번호 = 목록 순번(1부터)
    bindata = {b.get("Id"): b for b in root.iter("BINDATA")}

    def image_bytes(bin_item: str) -> tuple[bytes, str]:
        try:
            item = items[int(bin_item) - 1]
            b = bindata.get(item.get("BinData"))
            raw = base64.b64decode("".join(b.itertext()))
            if (b.get("Compress") or "").lower() == "true":
                try:
                    raw = zlib.decompress(raw, -15)
                except zlib.error:
                    raw = zlib.decompress(raw)
            return raw, item.get("Format", "")
        except Exception:
            return b"", ""

    blocks: list[str] = []
    n_tables = 0

    def text_of_char(ch) -> str:
        parts = [ch.text or ""]
        for sub in ch:
            if sub.tag == "TAB":
                parts.append("\t")
            elif sub.tag == "LINEBREAK":
                parts.append("\n")
            parts.append(sub.tail or "")
        return "".join(parts)

    def cell_text(cell) -> str:
        lines = []
        for p in cell.iter("P"):
            if p is cell:
                continue
            s = "".join(text_of_char(c) for t in p.findall("TEXT") for c in t.findall("CHAR"))
            inner = [cell_text(tc) for t in p.findall("TEXT") for tb in t.findall("TABLE")
                     for tc in tb.iter("CELL")]
            s = (s + (" / " + " / ".join(x for x in inner if x) if inner else "")).strip()
            if s:
                lines.append(s)
        return "\n".join(dict.fromkeys(lines))

    def table_md(tb) -> str:
        cells = {}
        for row in tb.findall("ROW"):
            for c in row.findall("CELL"):
                try:
                    r, col = int(c.get("RowAddr", 0)), int(c.get("ColAddr", 0))
                except ValueError:
                    continue
                if r <= 500 and col <= 200:
                    plist = c.find("PARALIST")
                    cells[(r, col)] = cell_text(plist if plist is not None else c)
        if not cells:
            return ""
        grid = [["" for _ in range(max(c for _, c in cells) + 1)]
                for _ in range(max(r for r, _ in cells) + 1)]
        for (r, col), v in cells.items():
            grid[r][col] = v
        return rows_to_md(grid)

    body = root.find("BODY")
    for sec in (body.findall("SECTION") if body is not None else []):
        for p in sec.findall("P"):
            buf = []
            for t in p.findall("TEXT"):
                for el in t:
                    if el.tag == "CHAR":
                        buf.append(text_of_char(el))
                    elif el.tag == "TABLE":
                        if "".join(buf).strip():
                            blocks.append("".join(buf).strip())
                        buf = []
                        md = table_md(el)
                        if md:
                            blocks.append(md)
                            n_tables += 1
                    elif el.tag == "PICTURE" and ctx.enabled:
                        img = el.find("IMAGE")
                        if img is not None and img.get("BinItem"):
                            data, fmt = image_bytes(img.get("BinItem"))
                            line = ctx.save_image(data, fmt)
                            if line:
                                blocks.append(line)
            if "".join(buf).strip():
                blocks.append("".join(buf).strip())
    if not blocks:
        return ReadResult(False, kind="한글(HML)", error="본문(BODY)에서 글을 찾지 못했습니다")
    return ReadResult(True, join_blocks(blocks), "한글(HML)", {"표": n_tables})


# ── 오픈도큐먼트 (.odt .ods .odp) ─────────────────────────────
_ODF = {
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "xlink": "http://www.w3.org/1999/xlink",
}


def _odf(tag: str) -> str:
    p, local = tag.split(":")
    return f"{{{_ODF[p]}}}{local}"


def _odf_inline(el) -> str:
    """문단 안의 글. <text:s c="3"/> 는 공백 3칸, <text:tab/> 은 탭이다."""
    parts = [el.text or ""]
    for ch in el:
        tag = ch.tag
        if tag == _odf("text:s"):
            parts.append(" " * int(ch.get(_odf("text:c"), "1") or 1))
        elif tag == _odf("text:tab"):
            parts.append("\t")
        elif tag == _odf("text:line-break"):
            parts.append("\n")
        elif tag in (_odf("text:note"), _odf("draw:frame")):
            pass                                            # 각주·그림은 따로 다룬다
        else:
            parts.append(_odf_inline(ch))
        parts.append(ch.tail or "")
    return "".join(parts)


def read_odf(path: str, ctx: ReadContext) -> ReadResult:
    kind = {"odt": "오픈도큐먼트 문서", "ods": "오픈도큐먼트 표", "odp": "오픈도큐먼트 발표"}
    try:
        z = zipfile.ZipFile(long_path(path))
    except Exception as e:
        return ReadResult(False, kind="오픈도큐먼트", error=f"열지 못했습니다: {e}")
    with z:
        names = set(z.namelist())
        if "content.xml" not in names:
            return ReadResult(False, kind="오픈도큐먼트", error="content.xml 이 없습니다")
        mt = z.read("mimetype").decode("ascii", "ignore") if "mimetype" in names else ""
        fmt = "odp" if "presentation" in mt else "ods" if "spreadsheet" in mt else "odt"
        root = ET.fromstring(z.read("content.xml"))
        blocks: list[str] = []
        n_tables = n_pages = 0

        def top_frames(el):
            """문단 안의 틀(draw:frame) 중 가장 바깥 것들. 안쪽 틀은 frame() 이 다룬다."""
            for ch in el:
                if ch.tag == _odf("draw:frame"):
                    yield ch
                else:
                    yield from top_frames(ch)

        def frame(fr, depth):
            """틀 하나 — 그림이거나, 글상자(안에 문단·표가 든다)다.
            한글에서 내보낸 ODT 는 표를 문단 안의 틀 속 글상자에 넣는다."""
            if ctx.enabled:
                for im in fr.findall(_odf("draw:image")):
                    href = im.get(_odf("xlink:href"), "")
                    if href in names:
                        line = ctx.save_image(z.read(href), os.path.splitext(href)[1])
                        if line:
                            blocks.append(line)
            walk(fr, depth + 1)

        def table(tb) -> str:
            rows = []
            for tr in tb.iter(_odf("table:table-row")):
                if tr.get(_odf("table:number-rows-repeated"), "1") not in ("1", "") and \
                        not "".join(tr.itertext()).strip():
                    continue                               # 시트 끝까지 반복되는 빈 행
                cells = []
                for tc in tr:
                    rep = int(tc.get(_odf("table:number-columns-repeated"), "1") or 1)
                    if tc.tag == _odf("table:covered-table-cell"):
                        cells += [""] * min(rep, 50)
                    elif tc.tag == _odf("table:table-cell"):
                        txt = "\n".join(_odf_inline(p).strip() for p in tc.iter(_odf("text:p")))
                        cells += [txt] * min(rep, 50)          # 같은 칸이 rep 번 이어진다
                rows.append(cells)
            return rows_to_md(rows)

        def walk(el, depth=0):
            nonlocal n_tables, n_pages
            for ch in el:
                tag = ch.tag
                if tag == _odf("text:h"):
                    lvl = int(ch.get(_odf("text:outline-level"), "1") or 1)
                    t = _odf_inline(ch).strip()
                    if t:
                        blocks.append("#" * min(max(lvl, 1), 6) + " " + t)
                elif tag == _odf("text:p"):
                    t = _odf_inline(ch).strip()
                    if t:
                        blocks.append(t)
                    for fr in top_frames(ch):
                        frame(fr, depth)
                elif tag == _odf("text:list"):
                    items = ["- " + _odf_inline(p).strip() for p in ch.iter(_odf("text:p"))
                             if _odf_inline(p).strip()]
                    if items:
                        blocks.append("\n".join(items))
                elif tag == _odf("table:table"):
                    if fmt == "ods":
                        blocks.append(f"## {ch.get(_odf('table:name'), '시트')}")
                    t = table(ch)
                    if t:
                        blocks.append(t)
                        n_tables += 1
                    # 칸 안의 그림은 표 바로 뒤에 둔다. 칸 안의 글상자는 따라가지 않는다 —
                    # 그 글은 table() 이 이미 칸 글로 옮겼으므로 두 번 적게 된다.
                    if ctx.enabled:
                        for im in ch.iter(_odf("draw:image")):
                            href = im.get(_odf("xlink:href"), "")
                            if href in names:
                                line = ctx.save_image(z.read(href), os.path.splitext(href)[1])
                                if line:
                                    blocks.append(line)
                elif tag == _odf("draw:page"):
                    n_pages += 1
                    blocks.append(f"## 슬라이드 {n_pages}")
                    walk(ch, depth + 1)
                elif tag == _odf("draw:frame"):
                    frame(ch, depth)
                elif depth < 12:
                    walk(ch, depth + 1)

        body = root.find(_odf("office:body"))
        walk(body if body is not None else root)
    return ReadResult(True, join_blocks(blocks), kind[fmt], {"표": n_tables})


# ═════════════════════════════════════════════════════════════
# 사진
# ═════════════════════════════════════════════════════════════
def read_image(path: str, kind, ctx: ReadContext) -> ReadResult:
    info = _img.describe_image(path)
    stem = os.path.splitext(os.path.basename(path))[0]
    meta = {k: v for k, v in info.items() if k in ("가로", "세로", "찍은날짜", "기기")}
    blocks = []

    ext = kind.ext if kind.ext else "." + ("jpg" if kind.fmt == "jpeg" else kind.fmt)
    link = ctx.save_original(path, "원본" + ext.lower())
    if link:
        blocks.append(md_link(stem, link, image=True))

    facts = []
    if info.get("찍은날짜"):
        facts.append(f"찍은 날짜 {info['찍은날짜']}")
    if info.get("제조사") or info.get("기기"):
        facts.append("기기 " + " ".join(x for x in (info.get("제조사"), info.get("기기")) if x))
    if info.get("가로"):
        facts.append(f"크기 {info['가로']}×{info['세로']}")
    if info.get("설명"):
        blocks.append(info["설명"])

    if kind.fmt == "svg":                             # 벡터 그림 — 안에 적힌 글자를 그대로 꺼낸다
        try:
            words = [t.strip() for t in ET.parse(long_path(path)).getroot().itertext() if t.strip()]
            if words:
                blocks.append(" ".join(words))
        except Exception:
            pass
    else:
        r = _ocr.run_ocr(path)
        if r.ok and r.text.strip():
            blocks.append(r.text.strip())
            meta["글자인식"] = r.engine
        elif _ocr.is_on():
            blocks.append(f"> 글자 인식({r.engine})으로 읽지 못했습니다 — {r.error}")
    if facts:
        blocks.append("> " + " · ".join(facts))

    # 이름이 IMG_0001 같으면 제목을 '사진 2019-07-21' 처럼 짓는다
    meta["제목후보"] = f"사진 {info['찍은날짜']}" if info.get("찍은날짜") else ""
    return ReadResult(True, join_blocks(blocks), kind.label, meta)


# ═════════════════════════════════════════════════════════════
# 읽을 수 없는 것 — 왜 못 읽는지, 어떻게 하면 되는지
# ═════════════════════════════════════════════════════════════
_CANNOT = {
    "empty":   "빈 파일입니다 (크기 0). 저장이 덜 되었거나 내려받다 끊긴 파일일 수 있습니다.",
    "hwp3":    "한글 97 이전 형식(HWP 3.0)입니다. 한글에서 열어 .hwpx 로 다시 저장한 뒤 넣어 주세요.",
    "encrypted_office": "암호가 걸린 오피스 문서입니다. 암호를 풀고 저장한 뒤 넣어 주세요.",
    "xlsb":    "엑셀 이진 통합문서(.xlsb)는 읽지 못합니다. 엑셀에서 .xlsx 로 저장해 주세요.",
    "gdoc":    "구글 문서는 내용이 온라인에만 있습니다. 드라이브에서 .docx 로 내려받아 넣어 주세요.",
    "gsheet":  "구글 시트는 내용이 온라인에만 있습니다. .xlsx 로 내려받아 넣어 주세요.",
    "gslides": "구글 슬라이드는 내용이 온라인에만 있습니다. .pptx 로 내려받아 넣어 주세요.",
    "epub":    "전자책은 아직 다루지 않습니다.",
    "rtf":     "서식글(RTF)은 워드에서 열어 .docx 로 저장해 주세요.",
    "7z":      "7z 압축은 파이썬 기본 기능으로 풀 수 없습니다. 직접 푼 뒤 넣어 주세요.",
    "rar":     "RAR 압축은 파이썬 기본 기능으로 풀 수 없습니다. 직접 푼 뒤 넣어 주세요.",
    "gzip":    "gzip 압축은 자동으로 풀지 않습니다. 직접 푼 뒤 넣어 주세요.",
    "bzip2":   "bzip2 압축은 자동으로 풀지 않습니다. 직접 푼 뒤 넣어 주세요.",
    "xz":      "xz 압축은 자동으로 풀지 않습니다. 직접 푼 뒤 넣어 주세요.",
    "tar":     "tar 묶음은 자동으로 풀지 않습니다. 직접 푼 뒤 넣어 주세요.",
    "zip":     "압축 파일입니다. --압축안열기 를 빼면 풀어서 안의 파일을 변환합니다.",
    "mp3":     "소리 파일은 글자가 없어 변환하지 않습니다. 자막(.srt)이 있으면 그것을 넣어 주세요.",
    "wav":     "소리 파일은 글자가 없어 변환하지 않습니다.",
    "ogg":     "소리 파일은 글자가 없어 변환하지 않습니다.",
    "flac":    "소리 파일은 글자가 없어 변환하지 않습니다.",
    "mp4":     "영상은 글자가 없어 변환하지 않습니다. 자막(.srt)이 있으면 그것을 넣어 주세요.",
    "exe":     "실행 파일은 자료가 아닙니다.",
    "sqlite":  "데이터베이스 파일은 다루지 않습니다.",
    "lnk":     "바로가기는 자료가 아닙니다. 가리키는 원본 파일을 넣어 주세요.",
}


def _cannot(kind, path: str = "") -> ReadResult:
    msg = _CANNOT.get(kind.fmt)
    if not msg:
        # '개설강좌.xlsx.pia' 처럼 문서 확장자 뒤에 낯선 확장자가 붙은 이진 파일 —
        # 보안(DRM)·암호화 프로그램이 원래 문서를 감싼 경우가 흔하다
        from .sniff import EXT_HINT, LABELS
        stem = os.path.splitext(os.path.basename(path))[0]
        inner = os.path.splitext(stem)[1].lower()
        if kind.fmt == "unknown" and inner in EXT_HINT:
            label = LABELS.get(EXT_HINT[inner], ("문서", ""))[0]
            msg = (f"원래 {label}({inner}) 파일을 다른 프로그램이 감싼 것으로 보입니다 "
                   f"(보안·DRM·암호화 프로그램 등). 그 프로그램에서 풀어 {inner} 로 저장한 뒤 "
                   f"넣어 주세요.")
        else:
            msg = ("무슨 형식인지 알아내지 못했습니다. 이름을 바꿔도 소용없으니, "
                   "원본을 만든 프로그램에서 다시 저장해 보세요.")
    return ReadResult(False, kind=kind.label, error=msg)


_TEXT_READERS = {
    "txt": read_plain, "md": read_markdown, "json": read_json, "srt": read_srt,
    "code": read_code, "xml": read_xml, "html": read_html,
}
_CTX_READERS = {
    "docx": read_docx, "pptx": read_pptx, "xlsx": read_xlsx, "pdf": read_pdf,
    "hwp": read_hwp, "hwpx": read_hwpx, "hml": read_hml,
    "odt": read_odf, "ods": read_odf, "odp": read_odf,
}
_PLAIN_READERS = {"doc": read_doc, "ppt": read_ppt, "xls": read_xls}
_IMAGE_FMTS = {"jpeg", "png", "gif", "webp", "bmp", "tiff", "heic", "ico", "svg"}


def read_by_kind(path: str, kind, ctx: ReadContext | None = None) -> ReadResult:
    """sniff 가 알아낸 FileKind 로 읽는다. 어떤 경우에도 예외를 던지지 않는다."""
    ctx = ctx or ReadContext()
    fmt = kind.fmt
    try:
        if fmt in _TEXT_READERS:
            res = _TEXT_READERS[fmt](path, kind.encoding)
        elif fmt in ("csv", "tsv"):
            res = read_delimited(path, kind.encoding, "\t" if fmt == "tsv" else ",")
        elif fmt in _CTX_READERS:
            res = _CTX_READERS[fmt](path, ctx)
        elif fmt in _PLAIN_READERS:
            res = _PLAIN_READERS[fmt](path)
        elif fmt in _IMAGE_FMTS:
            res = read_image(path, kind, ctx)
        else:
            res = _cannot(kind, path)
    except Exception as e:
        res = ReadResult(False, kind=kind.label,
                         error=f"읽는 중 오류 ({type(e).__name__}): {e}")
    if res.ok and kind.mismatch:
        res.meta = dict(res.meta or {})
        res.meta["형식불일치"] = f"{kind.ext} → {kind.fmt}"
    return res


# ═════════════════════════════════════════════════════════════
# 제목 알아내기 — 파일 이름이 엉망일 때를 위해
# ═════════════════════════════════════════════════════════════
_JUNK_STEM = (
    r"문서|제목\s*없음|무제|새\s*문서|새\s*폴더|이름\s*없음|"
    r"image|img|photo|pic|picture|pasted\s*image|screenshot|screen\s*shot|"
    r"화면\s*캡처|화면\s*갈무리|캡처|다운로드|download|untitled|document|"
    r"doc|scan|스캔|kakaotalk|카카오톡|received|copy|사본|복사본|new|dsc|dcim|pxl"
)
_JUNK_NAME = re.compile(rf"^\s*(?:{_JUNK_STEM})[\s_\-().]*[\d\s_\-().]*$", re.I)
_ONLY_NUMBERS = re.compile(r"^[\d\s_\-().:]+$")


def is_junk_name(stem: str) -> bool:
    """파일 이름이 내용을 전혀 알려주지 못하는가"""
    s = (stem or "").strip()
    if len(s) <= 1:
        return True
    return bool(_JUNK_NAME.match(s)) or bool(_ONLY_NUMBERS.match(s))


_NOT_TITLE = re.compile(
    r"^(슬라이드|slide|시트|sheet|쪽|페이지|page|표|table|그림|figure|"
    r"문서\s*속\s*그림|발표자\s*노트|머리글.*|각주|미주)\s*\d*\.?\s*$", re.I)


def guess_title(text: str, fallback: str, hint: str = "") -> tuple[str, str]:
    """→ (제목, 어디서 왔는지)

    파일 이름이 쓸 만하면 그대로 쓴다. 'Pasted image 20260815' 처럼 아무 정보가
    없을 때만 읽기 함수가 준 후보(hint) → 본문 첫 줄 순으로 찾는다.
    """
    if not is_junk_name(fallback):
        return fallback, "파일이름"
    if hint and hint.strip():
        return hint.strip()[:80], "문서 속 제목"
    for raw in (text or "").splitlines()[:60]:
        s = raw.strip()
        if not s or s.startswith(("---", "```", "|", ">", "!", "<!--", "- ", "* ")):
            continue
        s = re.sub(r"^#{1,6}\s*(?:\d+\.\s*)?", "", s)
        s = re.sub(r"[*_`]", "", s).strip()
        if len(s) < 2 or len(s) > 80 or _NOT_TITLE.match(s):
            continue
        if re.match(r"^[\d\s\-./:]+$", s):
            continue
        return s, "본문 첫 줄"
    return fallback, "파일이름"


def strip_leading_h1(text: str, title: str) -> str:
    """머리말이 이미 '# 제목' 을 넣으므로, 본문 맨 앞의 같은 제목 줄은 지운다."""
    lines = (text or "").split("\n")
    for i, ln in enumerate(lines[:5]):
        s = ln.strip()
        if not s:
            continue
        m = re.match(r"^#\s+(.+?)\s*$", s)
        if m and m.group(1).strip() == (title or "").strip():
            del lines[i]
            while i < len(lines) and not lines[i].strip():
                del lines[i]
        break
    return "\n".join(lines)
